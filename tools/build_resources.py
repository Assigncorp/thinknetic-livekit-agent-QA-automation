#!/usr/bin/env python3
"""
Regenerates a product's generated resources from its test-bed config.

Which config, and therefore which product, comes from TESTBED_CONFIG in the
repo-root .env (default: config/testbed.config.json) - the same variable the
browser and API suites read, so all three always describe the same product.

Two outputs, both consumed by the TypeScript and Python suites alike:

  serial-index.json  serial number -> knowledge base, derived from the hopper
                     classification workbook. This is what makes serial rotation
                     possible without hand-maintaining a list.

  scenarios.json     caller questions extracted from each KB markdown file,
                     with anchor terms pulled from the KB's own answer.

Run after changing a KB file, renaming one, or updating the workbook:

    make resources
"""

from __future__ import annotations

import json
import os
import random
import re
import sys
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: run `uv sync` in tools/, or `pip install openpyxl`")

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = "config/testbed.config.json"


def _load_env() -> None:
    """Read TESTBED_CONFIG from the repo-root .env without requiring dotenv."""
    env_file = ROOT / ".env"
    if not env_file.exists():
        return
    for line in env_file.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, value = line.partition("=")
        os.environ.setdefault(key.strip(), value.strip())


def load_config() -> dict[str, Any]:
    _load_env()
    path = ROOT / os.getenv("TESTBED_CONFIG", DEFAULT_CONFIG)
    if not path.exists():
        sys.exit(f"config not found: {path}  (TESTBED_CONFIG in .env)")
    print(f"Using {path.relative_to(ROOT)}")
    return json.loads(path.read_text(encoding="utf-8"))


# --------------------------------------------------------------------------
# Serial index
# --------------------------------------------------------------------------

def build_serial_index(cfg: dict[str, Any]) -> dict[str, Any]:
    workbook_path = ROOT / cfg["resources"]["serialWorkbook"]
    if not workbook_path.exists():
        sys.exit(f"serial workbook not found: {workbook_path}")

    # (controller family, hopper type) -> kb id
    routing: dict[tuple[str, str], str] = {}
    for kb in cfg["knowledgeBases"]:
        if kb.get("enabled") and kb.get("controller") and kb.get("hopperType"):
            routing[(kb["controller"], kb["hopperType"])] = kb["id"]

    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
    excluded = {s.upper() for s in cfg["scenarioSelection"].get("excludeSerials", [])}
    per_kb_cap = int(cfg["scenarioSelection"].get("serialsPerKb", 0))

    collected: dict[str, list[dict[str, str]]] = {}
    unroutable = 0

    for sheet_cfg in cfg["serialSource"]["sheets"]:
        name = sheet_cfg["name"]
        if name not in wb.sheetnames:
            print(f"  ! sheet '{name}' not in workbook - skipped")
            continue

        ws = wb[name]
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]

        def col(key: str) -> int:
            label = sheet_cfg[key]
            if label not in header:
                sys.exit(f"column '{label}' not found in sheet '{name}' (has: {header})")
            return header.index(label)

        i_serial, i_desc, i_hopper = col("serialColumn"), col("descriptionColumn"), col("hopperTypeColumn")
        family = sheet_cfg["family"]
        seen: set[str] = set()

        for row in rows:
            if row is None or len(row) <= max(i_serial, i_desc, i_hopper):
                continue
            serial = str(row[i_serial] or "").strip().upper()
            hopper = str(row[i_hopper] or "").strip().upper()
            if not serial or serial in seen or serial in excluded:
                continue
            seen.add(serial)

            kb_id = routing.get((family, hopper))
            if kb_id is None:
                unroutable += 1
                continue

            collected.setdefault(kb_id, []).append(
                {
                    "serial": serial,
                    "description": str(row[i_desc] or "").strip(),
                    "controller": family,
                    "hopperType": hopper,
                    "sheet": name,
                    "kbId": kb_id,
                }
            )

    # Keep the anchor serial first, then a deterministic sample of the rest, so
    # the index is stable across regenerations but still exercises the catalogue.
    anchors = {kb["id"]: kb.get("anchorSerial") for kb in cfg["knowledgeBases"]}
    entries: list[dict[str, str]] = []
    rng = random.Random(1337)

    for kb_id, items in collected.items():
        items.sort(key=lambda e: e["serial"])
        anchor = anchors.get(kb_id)
        pinned = [e for e in items if e["serial"] == anchor]
        rest = [e for e in items if e["serial"] != anchor]
        if per_kb_cap > 0:
            keep = max(0, per_kb_cap - len(pinned))
            rest = rng.sample(rest, min(keep, len(rest)))
            rest.sort(key=lambda e: e["serial"])
        entries.extend(pinned + rest)

    by_kb = {k: len([e for e in entries if e["kbId"] == k]) for k in sorted(collected)}
    print(f"  serials: {len(entries)} kept ({by_kb}); {unroutable} rows had no matching KB")

    return {
        "_generated": "Do not edit. Regenerate with: make resources",
        "countsByKb": by_kb,
        "serials": entries,
    }


# --------------------------------------------------------------------------
# Scenario pool
# --------------------------------------------------------------------------

def anchor_terms(answer: str, limit: int = 6, asked: str = "") -> list[str]:
    """
    Facts from the KB's own answer that a correct reply cannot paraphrase away:
    pin references, part numbers, and measurements with units.

    Deliberately returns nothing rather than padding with ordinary vocabulary.
    An anchor like "switch" or "display" appears in almost any plausible reply,
    so asserting on it would pass regardless of whether the agent answered from
    the right knowledge base - a check that cannot fail is worse than no check,
    because it reads as coverage. Scenarios with no anchors are simply not
    content-checked (see assertions.checkExpectedAnchors).

    Anything already in `asked` is dropped for the same reason: the caller reads
    the question aloud, so a fact that appears in it proves only that the agent
    echoed us back. "Machine speed is limited to 400 FPM" must not be evidence
    that the agent looked up 400 FPM.
    """
    terms: list[str] = []

    for pattern in (
        # Pin references and part numbers.
        r"\b[A-Z]{1,2}\d-PIN\s*\d+\b",
        r"\b\d{6,7}\b",
        # Measurements with a unit. Every one of these is a number the manual
        # commits to, which a correct answer has to get right.
        r"\b[\d,]+(?:\.\d+)?\s?(?:PSI|FPM|RPM|ohms?|volts?|VDC|amps?|gallons?)\b",
        r"\b\d+(?:\.\d+)?\s?°?\s?F\b",
        r"\b\d+(?:\.\d+)?\s?(?:inch|inches)\b",
        # Clearances are always written as a fraction WITH a unit - a bare
        # "20/21" is a figure reference, not a specification.
        r"\b\d+/\d+\s?(?:inch|inches|\")",
    ):
        terms.extend(m.group(0) for m in re.finditer(pattern, answer, re.IGNORECASE))

    # De-duplicate, preserve order, and drop anything the question already says.
    asked_lower = asked.lower()
    out: list[str] = []
    for t in terms:
        if t.lower() in {o.lower() for o in out}:
            continue
        if asked_lower and t.lower() in asked_lower:
            continue
        out.append(t)
    return out[:limit]


def collect_answer(lines: list[str], idx: int, stoppers: tuple, window: int = 30) -> str:
    """Text belonging to the entry at `idx`, up to the next heading or question."""
    out: list[str] = []
    for follow in lines[idx + 1: idx + 1 + window]:
        if follow.startswith("#"):
            break
        if any(p and p.match(follow) for p in stoppers):
            break
        stripped = follow.strip()
        if stripped in ("", "---"):
            continue
        out.append(stripped)
    return " ".join(out)


def build_scenarios(cfg: dict[str, Any]) -> dict[str, Any]:
    extraction = cfg["scenarioExtraction"]
    faq_re = re.compile(extraction["faqPattern"])
    howto_re = re.compile(extraction["howToPattern"])
    problem_re = re.compile(extraction["problemPattern"]) if extraction.get("problemPattern") else None
    symptom_re = re.compile(extraction["problemSymptomPattern"]) if extraction.get("problemSymptomPattern") else None
    min_len = int(extraction["minQuestionLength"])
    cap = int(extraction.get("maxScenariosPerKb", 0))

    kb_dir = ROOT / cfg["resources"]["kbDir"]
    pools: dict[str, list[dict[str, Any]]] = {}

    for kb in cfg["knowledgeBases"]:
        if not kb.get("enabled"):
            continue
        path = kb_dir / kb["file"]
        if not path.exists():
            sys.exit(f"KB file declared in config but missing on disk: {path}")

        lines = path.read_text(encoding="utf-8").splitlines()
        scenarios: list[dict[str, Any]] = []
        section = ""
        counter = 0

        for idx, line in enumerate(lines):
            if line.startswith("## "):
                section = line.lstrip("# ").strip()

            faq = faq_re.match(line)
            howto = howto_re.match(line)
            problem = problem_re.match(line) if problem_re else None
            if not (faq or howto or problem):
                continue

            if problem:
                # The caller describes a symptom, not a chapter title. Prefer the
                # entry's own **Symptoms:** line as the thing a caller would say.
                question = problem.group("question").strip()
                if symptom_re:
                    for follow in lines[idx + 1: idx + 8]:
                        m = symptom_re.match(follow)
                        if m:
                            question = m.group("symptom").strip().rstrip(".")
                            break
            else:
                question = (faq or howto).group("question").strip().rstrip("*").strip()
            if len(question) < min_len:
                continue

            # The answer is the text until the next question or heading.
            answer = collect_answer(lines, idx, (faq_re, howto_re, problem_re))
            if not answer:
                continue

            counter += 1
            prompt = question.lstrip("Q:").strip()
            scenarios.append(
                {
                    "id": f"{kb['id'].upper()}-{'FAQ' if faq else ('HOW' if howto else 'PRB')}-{counter:03d}",
                    "kbId": kb["id"],
                    "kind": "faq" if faq else ("howto" if howto else "problem"),
                    "section": section,
                    "question": prompt,
                    "expectAnchors": anchor_terms(answer, asked=prompt),
                    "sourceLine": idx + 1,
                }
            )

        if cap > 0:
            scenarios = scenarios[:cap]
        pools[kb["id"]] = scenarios
        print(f"  scenarios: {kb['id']:<8} {len(scenarios):>4}  ({path.name})")

    return {
        "_generated": "Do not edit. Regenerate with: make resources",
        "countsByKb": {k: len(v) for k, v in pools.items()},
        "pools": pools,
    }


def main() -> None:
    cfg = load_config()
    out_dir = ROOT / cfg["resources"]["generatedDir"]
    out_dir.mkdir(parents=True, exist_ok=True)

    print("Building scenario pool...")
    scenarios = build_scenarios(cfg)
    (out_dir / "scenarios.json").write_text(json.dumps(scenarios, indent=1) + "\n", encoding="utf-8")

    print("Building serial index...")
    index = build_serial_index(cfg)
    (out_dir / "serial-index.json").write_text(json.dumps(index, indent=1) + "\n", encoding="utf-8")

    print(f"\nWrote {out_dir.relative_to(ROOT)}/scenarios.json and serial-index.json")


if __name__ == "__main__":
    main()
